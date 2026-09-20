import csv
import json
from pathlib import Path
from typing import Any, Dict, Iterable, List
from urllib.parse import urlparse

import requests
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_ACTIVITY_FILE = PROJECT_ROOT / "data" / "activities.json"


def normalize_participants(raw_participants: Any) -> List[str]:
    """Normalize participant records from CSV/Excel/JSON into a clean list."""
    if raw_participants is None:
        return []

    if isinstance(raw_participants, list):
        values = raw_participants
    elif isinstance(raw_participants, str):
        values = [part.strip() for part in raw_participants.replace(";", ",").split(",")]
    else:
        values = [str(raw_participants)]

    normalized = []
    for value in values:
        if isinstance(value, str):
            item = value.strip()
        else:
            item = str(value).strip()
        if item:
            normalized.append(item)
    return normalized


def normalize_activity_record(raw_record: Dict[str, Any]) -> Dict[str, Any]:
    """Convert a row from an import source into the internal activity format."""
    if not isinstance(raw_record, dict):
        raise ValueError("Activity rows must be dictionaries")

    name = str(raw_record.get("name") or raw_record.get("activity_name") or raw_record.get("title") or "").strip()
    if not name:
        raise ValueError("Each activity row must contain a name")

    max_participants_raw = raw_record.get("max_participants")
    if max_participants_raw is None:
        max_participants_raw = raw_record.get("capacity")
    if max_participants_raw is None:
        max_participants_raw = raw_record.get("maxParticipants")

    try:
        max_participants = int(max_participants_raw)
    except (TypeError, ValueError):
        max_participants = 0

    return {
        "description": str(raw_record.get("description") or "").strip(),
        "schedule": str(raw_record.get("schedule") or "").strip(),
        "max_participants": max_participants,
        "participants": normalize_participants(raw_record.get("participants") or raw_record.get("student_emails") or []),
        "name": name,
    }


def import_activity_rows(rows: Iterable[Dict[str, Any]]) -> Dict[str, Dict[str, Any]]:
    """Convert a sequence of rows into a name-keyed dictionary of activities."""
    activities: Dict[str, Dict[str, Any]] = {}
    for row in rows:
        if row is None:
            continue
        normalized = normalize_activity_record(row)
        name = normalized.pop("name")
        activities[name] = {
            "description": normalized["description"],
            "schedule": normalized["schedule"],
            "max_participants": normalized["max_participants"],
            "participants": normalized["participants"],
        }
    return activities


def import_activity_file(source: str | Path) -> Dict[str, Dict[str, Any]]:
    """Import activities from JSON, CSV, XLSX, or a Google Sheets export URL."""
    source_path = str(source)

    if source_path.startswith("http://") or source_path.startswith("https://"):
        return import_google_sheet(source_path)

    path = Path(source_path)
    if not path.exists():
        raise FileNotFoundError(f"Activity source not found: {path}")

    suffix = path.suffix.lower()

    if suffix == ".json":
        payload = json.loads(path.read_text(encoding="utf-8"))
        if isinstance(payload, dict):
            normalized: Dict[str, Dict[str, Any]] = {}
            for name, details in payload.items():
                if isinstance(details, dict):
                    normalized[name] = {
                        "description": str(details.get("description") or "").strip(),
                        "schedule": str(details.get("schedule") or "").strip(),
                        "max_participants": int(details.get("max_participants", 0) or 0),
                        "participants": normalize_participants(details.get("participants") or []),
                    }
            return normalized
        return import_activity_rows(payload)

    if suffix == ".csv":
        with path.open("r", encoding="utf-8", newline="") as handle:
            reader = csv.DictReader(handle)
            return import_activity_rows(list(reader))

    if suffix in {".xlsx", ".xls"}:
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return {}

        headers = [str(value).strip() for value in rows[0]]
        data_rows = []
        for row in rows[1:]:
            record = {}
            for idx, header in enumerate(headers):
                if idx < len(row):
                    record[header] = row[idx]
            if record and any(value not in (None, "") for value in record.values()):
                data_rows.append(record)
        workbook.close()
        return import_activity_rows(data_rows)

    raise ValueError(f"Unsupported activity import format: {suffix}")


def import_google_sheet(url: str) -> Dict[str, Dict[str, Any]]:
    """Fetch a Google Sheets CSV export and import it into the app format."""
    parsed = urlparse(url)
    if "docs.google.com" not in parsed.netloc and "drive.google.com" not in parsed.netloc:
        raise ValueError("Only Google Sheets or Drive export URLs are supported")

    export_url = url
    if "export?format=csv" not in export_url and "/pub?output=csv" not in export_url:
        if "/spreadsheets/d/" in export_url:
            export_url = export_url.replace("/edit", "/export?format=csv")
        elif "/file/d/" in export_url:
            export_url = export_url.replace("/view?usp=sharing", "/export?format=csv")

    response = requests.get(export_url, timeout=30)
    response.raise_for_status()

    rows = list(csv.DictReader(response.text.splitlines()))
    return import_activity_rows(rows)


def load_activities(path: str | Path | None = None) -> Dict[str, Dict[str, Any]]:
    """Load activities from the JSON file, creating it if needed."""
    target = Path(path) if path else DEFAULT_ACTIVITY_FILE
    target.parent.mkdir(parents=True, exist_ok=True)

    if not target.exists():
        default_activities = {
            "Chess Club": {
                "description": "Learn strategies and compete in chess tournaments",
                "schedule": "Fridays, 3:30 PM - 5:00 PM",
                "max_participants": 12,
                "participants": ["michael@mergington.edu", "daniel@mergington.edu"],
            },
            "Programming Class": {
                "description": "Learn programming fundamentals and build software projects",
                "schedule": "Tuesdays and Thursdays, 3:30 PM - 4:30 PM",
                "max_participants": 20,
                "participants": ["emma@mergington.edu", "sophia@mergington.edu"],
            },
            "Gym Class": {
                "description": "Physical education and sports activities",
                "schedule": "Mondays, Wednesdays, Fridays, 2:00 PM - 3:00 PM",
                "max_participants": 30,
                "participants": ["john@mergington.edu", "olivia@mergington.edu"],
            },
            "Soccer Team": {
                "description": "Join the school soccer team and compete in matches",
                "schedule": "Tuesdays and Thursdays, 4:00 PM - 5:30 PM",
                "max_participants": 22,
                "participants": ["liam@mergington.edu", "noah@mergington.edu"],
            },
            "Basketball Team": {
                "description": "Practice and play basketball with the school team",
                "schedule": "Wednesdays and Fridays, 3:30 PM - 5:00 PM",
                "max_participants": 15,
                "participants": ["ava@mergington.edu", "mia@mergington.edu"],
            },
            "Art Club": {
                "description": "Explore your creativity through painting and drawing",
                "schedule": "Thursdays, 3:30 PM - 5:00 PM",
                "max_participants": 15,
                "participants": ["amelia@mergington.edu", "harper@mergington.edu"],
            },
            "Drama Club": {
                "description": "Act, direct, and produce plays and performances",
                "schedule": "Mondays and Wednesdays, 4:00 PM - 5:30 PM",
                "max_participants": 20,
                "participants": ["ella@mergington.edu", "scarlett@mergington.edu"],
            },
            "Math Club": {
                "description": "Solve challenging problems and participate in math competitions",
                "schedule": "Tuesdays, 3:30 PM - 4:30 PM",
                "max_participants": 10,
                "participants": ["james@mergington.edu", "benjamin@mergington.edu"],
            },
            "Debate Team": {
                "description": "Develop public speaking and argumentation skills",
                "schedule": "Fridays, 4:00 PM - 5:30 PM",
                "max_participants": 12,
                "participants": ["charlotte@mergington.edu", "henry@mergington.edu"],
            },
        }
        save_activities(default_activities, target)
        return default_activities

    payload = json.loads(target.read_text(encoding="utf-8"))
    if isinstance(payload, dict):
        normalized: Dict[str, Dict[str, Any]] = {}
        for name, details in payload.items():
            if not isinstance(details, dict):
                continue
            normalized[name] = {
                "description": str(details.get("description") or "").strip(),
                "schedule": str(details.get("schedule") or "").strip(),
                "max_participants": int(details.get("max_participants", 0) or 0),
                "participants": normalize_participants(details.get("participants") or []),
            }
        return normalized

    return import_activity_rows(payload)


def save_activities(activities: Dict[str, Dict[str, Any]], path: str | Path | None = None) -> None:
    """Persist activities to JSON."""
    target = Path(path) if path else DEFAULT_ACTIVITY_FILE
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(activities, ensure_ascii=False, indent=2), encoding="utf-8")


__all__ = [
    "DEFAULT_ACTIVITY_FILE",
    "import_activity_rows",
    "import_activity_file",
    "import_google_sheet",
    "load_activities",
    "save_activities",
    "normalize_activity_record",
]
